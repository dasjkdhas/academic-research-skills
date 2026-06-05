"""Apply 3 targeted edits to the user's uploaded latest pptx (20260601).
Preserves all other content/structure (notes, depth pages 12-14, etc.)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import shutil, os, re

SRC = '/root/.claude/uploads/52631f48-0c58-4d9a-a61c-63bae5283405/78519c36-HIRAKU_Global_Interview_Liu_20260601.pptx'
DST = '/home/user/academic-research-skills/interview_materials/HIRAKU_Global_Interview_Liu.pptx'
BAK = '/home/user/academic-research-skills/interview_materials/HIRAKU_Global_Interview_Liu_pre_v3_backup.pptx'

# Backup current main file
if os.path.exists(DST):
    shutil.copy2(DST, BAK)
    print(f'Backed up previous main → {BAK}')

p = Presentation(SRC)

# ===== EDIT 1: Slide 10 — Teaching & Mentoring card body =====
print('\n[Edit 1] Slide 10 · TEACHING & MENTORING card')
s10 = p.slides[9]
done1 = False
for sh in s10.shapes:
    if not sh.has_text_frame: continue
    if sh.text_frame.text.strip() == 'Grad supervision · seminars':
        # Capture original styling (color from first run)
        orig_run = sh.text_frame.paragraphs[0].runs[0]
        orig_color = None
        try:
            orig_color = orig_run.font.color.rgb
        except Exception:
            pass
        # Extend shape height to fit two lines
        old_h = sh.height
        sh.height = Inches(0.36)
        # Clear and rewrite
        tf = sh.text_frame
        tf.clear()
        # Line 1: bold heading
        p0 = tf.paragraphs[0]
        r0 = p0.add_run()
        r0.text = 'Mentored 3 PhDs'
        r0.font.name = 'Calibri'
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.italic = False
        if orig_color:
            r0.font.color.rgb = orig_color
        # Line 2: italic detail
        p1 = tf.add_paragraph()
        r1 = p1.add_run()
        r1.text = 'Tsinghua D3 · Okayama D2 · U-Tokyo (admitted)'
        r1.font.name = 'Calibri'
        r1.font.size = Pt(7)
        r1.font.bold = False
        r1.font.italic = True
        if orig_color:
            r1.font.color.rgb = orig_color
        print(f'  ✓ Rewrote shape (h: {old_h/914400:.2f}" → 0.36")')
        done1 = True
        break
if not done1:
    print('  ✗ Target shape NOT FOUND')

# ===== EDIT 2: Slide 6 — Notes (honest authorship + 3 students) =====
print('\n[Edit 2] Slide 6 · NOTES rewrite')
s6 = p.slides[5]
new_notes_6 = (
    "I have 15 peer-reviewed papers, including 10 in English international "
    "journals. Three are in top-tier — Q1 — journals.\n\n"
    "I want to be precise on authorship. In two of those three I was first "
    "or co-first author. In the third I was a team contributor.\n\n"
    "The lead student on the Building and Environment paper is one of three "
    "PhD students I have mentored — now at Tsinghua, Okayama, and the "
    "University of Tokyo.\n\n"
    "The topics form one line — building, health, air, urban heat, and policy."
)
s6.notes_slide.notes_text_frame.text = new_notes_6
print(f'  ✓ Replaced notes ({len(new_notes_6.split())} spoken words)')

# ===== EDIT 3: Slide 11 — Prepend timing prompt to notes =====
print('\n[Edit 3] Slide 11 · NOTES prepend timing prompt')
s11 = p.slides[10]
existing_11 = s11.notes_slide.notes_text_frame.text
prompt_11 = (
    "[TIME CHECK ~8:00 — if under 1:30 remains, skip 12-14 "
    "and jump to Slide 15.]\n\n"
)
s11.notes_slide.notes_text_frame.text = prompt_11 + existing_11
print(f'  ✓ Prepended timing prompt to existing notes')

# Save
p.save(DST)
print(f'\n✅ Saved: {DST}')

# Verify timing
print('\n=== TIMING AUDIT ===')
p2 = Presentation(DST)
total_notes = 0
total_spoken = 0
for i, sl in enumerate(p2.slides, 1):
    notes = sl.notes_slide.notes_text_frame.text
    spoken = re.sub(r'\[[^\]]+\]', '', notes)  # drop [stage directions]
    w = len(notes.split())
    sw = len(spoken.split())
    total_notes += w
    total_spoken += sw
    tag = ''
    if i == 6: tag = ' ← edited'
    if i == 10: tag = ' ← (card edited, notes unchanged)'
    if i == 11: tag = ' ← stage direction added'
    print(f'  Slide {i:2d}: notes={w:3d}w  spoken={sw:3d}w{tag}')
print(f'\nTotal spoken words: {total_spoken}')
print(f'  @100 wpm (safe):   {total_spoken*60/100/60:.2f} min = {int(total_spoken*60/100)}s')
print(f'  @110 wpm (normal): {total_spoken*60/110/60:.2f} min = {int(total_spoken*60/110)}s')
print(f'  @120 wpm (fluent): {total_spoken*60/120/60:.2f} min = {int(total_spoken*60/120)}s')
print(f'\n+ 15 page transitions (~5s each) ≈ +1:15')
print(f'  Expected total at safe pace: ~{int(total_spoken*60/100) + 75}s = {(total_spoken*60/100 + 75)/60:.2f} min')
