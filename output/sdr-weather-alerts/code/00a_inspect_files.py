"""Step 0a: Inspect raw Excel structure of all 5 uploaded files."""
import os
import warnings
warnings.filterwarnings('ignore')

import pandas as pd
from openpyxl import load_workbook

DATA_DIR = "/home/user/academic-research-skills/output/sdr-weather-alerts/data_raw"

files = sorted([f for f in os.listdir(DATA_DIR) if f.endswith(('.xlsm', '.xlsx'))])

print("=" * 80)
print(f"Found {len(files)} files")
print("=" * 80)

for fname in files:
    fpath = os.path.join(DATA_DIR, fname)
    print(f"\n{'#'*80}")
    print(f"# FILE: {fname}")
    print(f"# Size: {os.path.getsize(fpath)/1024:.1f} KB")
    print(f"{'#'*80}")
    try:
        wb = load_workbook(fpath, read_only=True, data_only=True)
        print(f"Sheets: {wb.sheetnames}")
        for sname in wb.sheetnames[:5]:
            ws = wb[sname]
            print(f"\n  --- Sheet: '{sname}'  (dim: {ws.max_row} rows x {ws.max_column} cols) ---")
            for i, row in enumerate(ws.iter_rows(values_only=True, max_row=12)):
                print(f"    R{i+1:02d}: {row[:10]}")
        wb.close()
    except Exception as e:
        print(f"ERROR: {e}")
